"""Tests for Stage 2: Policy Parser."""

import pytest
import tempfile
from pathlib import Path
from docx import Document
import yaml

from src.expense_pipeline.policy_parser import (
    parse_policy_manual,
    load_policy_rules_yaml,
)
from src.expense_pipeline.gates import gate_policy_parser, GateFailure


def create_test_policy_doc(path: str):
    """Create a test policy manual document."""
    doc = Document()

    doc.add_heading("Company Expense Policy", level=1)
    doc.add_paragraph("This is our expense policy.")

    doc.add_heading("Section 4: Category Limits", level=2)

    # Add a table with policy rules
    table = doc.add_table(rows=7, cols=4)
    table.style = "Light Grid Accent 1"

    # Header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Category"
    hdr_cells[1].text = "Daily Limit"
    hdr_cells[2].text = "Receipt Required Above"
    hdr_cells[3].text = "Manager Approval Above"

    # Data rows
    categories_data = [
        ("meals", "75", "30", ""),
        ("travel", "500", "150", ""),
        ("lodging", "300", "100", ""),
        ("software", "1000", "500", "750"),
        ("office_supplies", "100", "50", ""),
        ("client_entertainment", "500", "200", "300"),
    ]

    for i, (category, limit, receipt, approval) in enumerate(categories_data, 1):
        row_cells = table.rows[i].cells
        row_cells[0].text = category
        row_cells[1].text = limit
        row_cells[2].text = receipt
        row_cells[3].text = approval

    doc.save(path)


def test_parse_policy_manual():
    """Test parsing a policy manual."""
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = Path(tmpdir) / "policy.docx"
        yaml_path = Path(tmpdir) / "policy_rules.yaml"

        create_test_policy_doc(str(doc_path))

        policy_rules = parse_policy_manual(str(doc_path), str(yaml_path))

        assert len(policy_rules.rules) >= 4
        assert "meals" in policy_rules.rules
        assert "travel" in policy_rules.rules
        assert "software" in policy_rules.rules

        meals_rule = policy_rules.rules["meals"]
        assert meals_rule.daily_limit == 75.0
        assert meals_rule.receipt_required_above == 30.0

        software_rule = policy_rules.rules["software"]
        assert software_rule.requires_manager_approval_above == 750.0

        # Check YAML was written
        assert yaml_path.exists()


def test_load_policy_rules_yaml():
    """Test loading policy rules from YAML."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yaml_path = Path(tmpdir) / "policy_rules.yaml"

        yaml_data = {
            "meals": {
                "daily_limit": 75,
                "receipt_required_above": 30,
            },
            "travel": {
                "daily_limit": 500,
                "receipt_required_above": 150,
            },
            "software": {
                "daily_limit": 1000,
                "receipt_required_above": 500,
                "requires_manager_approval_above": 750,
            },
        }

        with open(yaml_path, "w") as f:
            yaml.dump(yaml_data, f)

        policy_rules = load_policy_rules_yaml(str(yaml_path))

        assert len(policy_rules.rules) == 3
        assert policy_rules.rules["meals"].daily_limit == 75
        assert policy_rules.rules["software"].requires_manager_approval_above == 750


def test_gate_policy_parser_all_categories_covered():
    """Test gate checks that all expense categories have policy rules."""
    policy_rules_data = {
        "meals": {"daily_limit": 75, "receipt_required_above": 30},
        "travel": {"daily_limit": 500, "receipt_required_above": 150},
    }

    from src.expense_pipeline.schemas import PolicyRule, PolicyRules

    rules = {
        cat: PolicyRule(
            category=cat,
            daily_limit=data["daily_limit"],
            receipt_required_above=data["receipt_required_above"],
        )
        for cat, data in policy_rules_data.items()
    }

    policy_rules = PolicyRules(rules=rules, source_file="test.yaml")

    # Test with categories that are covered
    expense_categories = {"meals", "travel"}
    gate_policy_parser(policy_rules, expense_categories)  # Should not raise

    # Test with missing category
    expense_categories_with_missing = {"meals", "travel", "software"}
    with pytest.raises(GateFailure):
        gate_policy_parser(policy_rules, expense_categories_with_missing)


def test_parse_policy_manual_file_not_found():
    """Test error when policy file doesn't exist."""
    with pytest.raises(FileNotFoundError):
        parse_policy_manual("/nonexistent/policy.docx")
